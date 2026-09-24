"""YAKE dynamic rule generation (plan 260924-1642-yake-dynamic-rules-doc-sync phase 01).

Text -> YAKE keywords -> spaCy EntityRuler JSON ``{"patterns": [{"label", "pattern"}]}``.
Ported from the yake_vi prototype (``/Users/hieplq1.aip/test/Yake``); this copy is
self-contained and never imports graphrag_ingest_langextract so it stays out of the
heavy sentence-transformers/gliner dependency chain. The ingestor calls
``ensure_rule_file`` on the text it has already loaded (no second document parse);
``load_text``/the CLI exist for manual runs and for dev.py's prune call only.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from numbers import Real
from pathlib import Path
from typing import Any, Iterable

import yake

RULE_FILE_PREFIX = "ruler.from-yake"

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b-\x1f\x7f-\x9f]")
_MULTI_SPACE = re.compile(r"\s+")


def clean_text(text: str) -> str:
    """NFC-normalize, strip control chars and collapse whitespace before extraction."""
    if not text:
        return ""
    normalized = unicodedata.normalize("NFC", text)
    normalized = _CONTROL_CHARS.sub(" ", normalized)
    normalized = _MULTI_SPACE.sub(" ", normalized)
    return normalized.strip()


@dataclass
class YakeRuleExtractor:
    """Thin wrapper around YAKE; defaults follow the harness doc pipeline (D6).

    ``dedup_lower=False`` keeps the prototype's ``dedupLim=False`` default.
    """

    language: str = "en"
    max_ngram: int = 3
    top: int = 150
    dedup_lower: bool = False

    def __post_init__(self) -> None:
        if self.max_ngram < 1:
            raise ValueError("max_ngram must be >= 1")
        if self.top < 1:
            raise ValueError("top must be >= 1")

        self._engine = yake.KeywordExtractor(
            lan=self.language,
            n=self.max_ngram,
            top=self.top,
            dedupLim=self.dedup_lower,
        )

    def extract(self, text: str, *, preprocess: bool = True) -> list[tuple[str, float]]:
        """Return ``(keyword, score)`` pairs sorted by score ascending (lower = better)."""
        payload = clean_text(text) if preprocess else (text or "").strip()
        if not payload:
            return []

        raw = self._engine.extract_keywords(payload)
        keywords = [(kw, float(score)) for kw, score in raw]
        keywords.sort(key=lambda item: item[1])
        return keywords


def infer_entity_label(phrase: str, *, default: str = "KEYWORD") -> str:
    """Best-effort label guess aligned with cortex-harness ruler.combined.json."""
    text = phrase.strip()
    lower = text.lower()

    if re.search(r"\brfc\s*\d+", lower):
        return "STANDARD"
    if re.search(r"\bieee\b", lower) or re.search(r"\b802\.\d+", lower):
        return "STANDARD"
    if re.match(r"iso[\s-]*\d+", lower):
        return "STANDARD"

    if re.search(r"\bnfc\b", lower) or "bluetooth" in lower or re.search(r"\buwb\b", lower):
        return "TRANSPORT"

    if re.search(r"\baes\b", lower) or re.search(r"\becdsa\b", lower) or re.search(r"\brsa\b", lower):
        return "CRYPTO"
    if re.search(r"\bhmac\b", lower) or re.search(r"\bhkdf\b", lower) or re.search(r"\bsha-?\d+", lower):
        return "KDF_HASH"
    if re.search(r"\bx\.509\b", lower) or re.search(r"\bca\b", lower) and "certificate" in lower:
        return "CERT"

    if "technical specification" in lower or lower.endswith(" specification"):
        return "DOCUMENT"
    if re.search(r"\bspake2\+?\b", lower) or "diffie-hellman" in lower or "curve25519" in lower:
        return "ALGORITHM"

    return default


def keywords_to_ruler_patterns(
    keywords: list[tuple[str, float]],
    *,
    default_label: str = "KEYWORD",
    use_heuristics: bool = True,
    min_length: int = 2,
) -> list[dict[str, Any]]:
    patterns: list[dict[str, Any]] = []
    seen: set[str] = set()

    for phrase, _score in keywords:
        cleaned = phrase.strip()
        if len(cleaned) < min_length:
            continue
        key = cleaned.casefold()
        if key in seen:
            continue
        seen.add(key)

        label = infer_entity_label(cleaned, default=default_label) if use_heuristics else default_label
        patterns.append({"label": label, "pattern": cleaned})

    return patterns


def build_ruler_json(
    keywords: list[tuple[str, float]],
    *,
    default_label: str = "KEYWORD",
    use_heuristics: bool = True,
    min_length: int = 2,
) -> dict[str, list[dict[str, Any]]]:
    return {
        "patterns": keywords_to_ruler_patterns(
            keywords,
            default_label=default_label,
            use_heuristics=use_heuristics,
            min_length=min_length,
        )
    }


def load_text(path: Path) -> str:
    """Read plain text for manual/CLI runs. NOT used by the ingestor hot path (D1).

    ``.xlsx`` returns "" (structured pipeline owns spreadsheets); PDF/docx/pptx
    readers are lazy imports so the module loads without them.
    """
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return ""
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n\n".join(p for p in parts if p.strip())
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        return "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text and p.text.strip())
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(str(path))
        parts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text and text.strip():
                    parts.append(text.strip())
        return "\n\n".join(parts)
    return path.read_text(encoding="utf-8")


def safe_source_key(source_id: str) -> str:
    """Sanitize a source id into a rule-file name component.

    Mirrors the ingestor's ``_safe_source_id`` separator rule (``/`` and ``\\``
    become ``__``) so relpath-based ids like ``docs/a/spec.pdf`` and
    ``docs/b/spec.pdf`` never collide (plan C3).
    """
    return str(source_id).replace("/", "__").replace("\\", "__")


def rule_file_for(rules_dir: Path | str, source_id: str) -> Path:
    return Path(rules_dir) / f"{RULE_FILE_PREFIX}.{safe_source_key(source_id)}.json"


def ensure_rule_file(
    rules_dir: Path | str,
    source_id: str,
    text: str,
    *,
    language: str = "en",
    top: int = 150,
    max_ngram: int = 3,
) -> Path | None:
    """Extract keywords from already-loaded ``text`` and persist the rule file.

    Empty pattern sets are never written; a stale rule file for this source is
    removed instead (plan C2 — empty files poison pattern-dir loaders).
    Returns the written path, or None when nothing was written.
    """
    extractor = YakeRuleExtractor(language=language, max_ngram=max_ngram, top=top)
    payload = build_ruler_json(extractor.extract(text))
    patterns = payload["patterns"]
    target = rule_file_for(rules_dir, source_id)
    if not patterns:
        if target.exists():
            target.unlink()
        return None
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=4) + "\n", encoding="utf-8")
    return target


def _iter_rule_files(rules_dir: Path | str) -> list[Path]:
    directory = Path(rules_dir)
    if not directory.is_dir():
        return []
    return sorted(directory.glob(f"{RULE_FILE_PREFIX}.*.json"))


def prune_rule_files(rules_dir: Path | str, keep_source_ids: Iterable[str]) -> list[str]:
    """Folder/full-sync prune: drop rule files whose source is not in the keep set."""
    keep_keys = {safe_source_key(source_id) for source_id in keep_source_ids}
    removed: list[str] = []
    for file_path in _iter_rule_files(rules_dir):
        source_key = _source_key_of(file_path)
        if source_key is not None and source_key not in keep_keys:
            file_path.unlink()
            removed.append(file_path.name)
    return removed


def prune_sources(rules_dir: Path | str, source_ids: Iterable[str]) -> list[str]:
    """Delete the rule files of explicitly deleted sources (dev.py D8 path)."""
    wanted = {safe_source_key(source_id) for source_id in source_ids}
    removed: list[str] = []
    for file_path in _iter_rule_files(rules_dir):
        source_key = _source_key_of(file_path)
        if source_key is not None and source_key in wanted:
            file_path.unlink()
            removed.append(file_path.name)
    return removed


def _source_key_of(file_path: Path) -> str | None:
    name = file_path.name
    if not name.startswith(f"{RULE_FILE_PREFIX}.") or not name.endswith(".json"):
        return None
    core = name[len(RULE_FILE_PREFIX) + 1 : -len(".json")]
    return core or None


def merge_rule_files(paths: Iterable[Path | str]) -> dict[str, list[dict[str, Any]]]:
    """Union patterns from several rule files, dedup on ``(label, pattern.casefold())``."""
    patterns: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for path in paths:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        for item in data.get("patterns", []) if isinstance(data, dict) else []:
            key = (str(item.get("label", "")), str(item.get("pattern", "")).casefold())
            if key in seen:
                continue
            seen.add(key)
            patterns.append({"label": item.get("label"), "pattern": item.get("pattern")})
    return {"patterns": patterns}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate spaCy EntityRuler patterns from document keywords via YAKE.",
    )
    parser.add_argument("--file", "-f", type=Path, help="Document to analyze (default: read stdin).")
    parser.add_argument("--top", "-t", type=int, default=150, help="Number of keywords (default: 150).")
    parser.add_argument("--n", type=int, default=3, dest="max_ngram", help="Max n-gram length (default: 3).")
    parser.add_argument("--language", "-l", default="en", metavar="LAN",
                        help="YAKE stopword language (default: en; use vi for Vietnamese docs).")
    parser.add_argument("--format", choices=("text", "ruler"), default="ruler",
                        help="Output format: EntityRuler JSON (default) or human text.")
    parser.add_argument("-o", "--output", type=Path, help="Write output to file (JSON for --format ruler).")
    parser.add_argument("--ruler-label", default="KEYWORD",
                        help="Default EntityRuler label when heuristics do not match (default: KEYWORD).")
    parser.add_argument("--no-ruler-heuristics", action="store_true",
                        help="Use --ruler-label for every pattern (no STANDARD/CRYPTO/... guessing).")
    parser.add_argument("--rules-dir", type=Path,
                        help="Rules directory for --prune/--prune-sources, and the default "
                             "output location for generated rule files (per-source naming).")
    parser.add_argument("--prune", nargs="+", metavar="KEEP_ID",
                        help="Delete rule files in --rules-dir except those for KEEP_ID sources, then exit.")
    parser.add_argument("--prune-sources", nargs="+", metavar="ID",
                        help="Delete rule files in --rules-dir for the given source ids, then exit.")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)

    if args.prune is not None or args.prune_sources is not None:
        if not args.rules_dir:
            print("--prune/--prune-sources require --rules-dir", file=sys.stderr)
            return 2
        if args.prune is not None:
            removed = prune_rule_files(args.rules_dir, args.prune)
            print(f"pruned {len(removed)} rule file(s) in {args.rules_dir} (keep-list mode)")
        else:
            removed = prune_sources(args.rules_dir, args.prune_sources)
            print(f"pruned {len(removed)} rule file(s) in {args.rules_dir} (deleted-source mode)")
        for name in removed:
            print(f"  - {name}")
        return 0

    text = load_text(args.file) if args.file is not None else sys.stdin.read()
    extractor = YakeRuleExtractor(
        language=args.language,
        max_ngram=args.max_ngram,
        top=args.top,
    )
    keywords = extractor.extract(text)

    if args.format == "ruler":
        payload = build_ruler_json(
            keywords,
            default_label=args.ruler_label,
            use_heuristics=not args.no_ruler_heuristics,
        )
        if not payload["patterns"]:
            # never write/emit an empty patterns payload — empty rule files
            # crash pattern-dir consumers (plan C2)
            print("0 patterns — nothing written")
            return 0
        body = json.dumps(payload, ensure_ascii=False, indent=4) + "\n"
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(body, encoding="utf-8")
        elif args.rules_dir:
            target = rule_file_for(args.rules_dir, args.file.stem if args.file else "stdin")
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(body, encoding="utf-8")
            print(f"{len(payload['patterns'])} patterns -> {target}")
        else:
            sys.stdout.write(body)
        return 0

    lines = [
        "Extracted keywords:",
        *(
            f"- {kw:<40} | score (lower is better): {score:.4f}"
            for kw, score in keywords
        ),
    ]
    text_out = "\n".join(lines) + "\n"
    if args.output:
        args.output.write_text(text_out, encoding="utf-8")
    else:
        print(text_out, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
